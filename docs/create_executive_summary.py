from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = 'Executive_Summary_Roadmap_Network_Phase_2-4.docx'
NAVY = '0B2545'
BLUE = '2E74B5'
LIGHT_BLUE = 'E8EEF5'
LIGHT_GRAY = 'F2F4F7'
GRAY = '5B6573'
GREEN = '1F6B45'
GOLD = '9A6A00'


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in('w:tcMar')
    if tc_mar is None:
        tc_mar = OxmlElement('w:tcMar')
        tc_pr.append(tc_mar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tc_mar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tc_mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_widths(table, widths_dxa):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn('w:tblW'))
    if tbl_w is None:
        tbl_w = OxmlElement('w:tblW')
        tbl_pr.append(tbl_w)
    tbl_w.set(qn('w:w'), str(sum(widths_dxa)))
    tbl_w.set(qn('w:type'), 'dxa')
    tbl_ind = tbl_pr.find(qn('w:tblInd'))
    if tbl_ind is None:
        tbl_ind = OxmlElement('w:tblInd')
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn('w:w'), '120')
    tbl_ind.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn('w:tcW'))
            if tc_w is None:
                tc_w = OxmlElement('w:tcW')
                tc_pr.append(tc_w)
            tc_w.set(qn('w:w'), str(widths_dxa[idx]))
            tc_w.set(qn('w:type'), 'dxa')
            set_cell_margins(cell)


def set_font(run, name='Calibri', size=11, color='000000', bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def add_para(doc, text='', style=None, size=11, color='000000', bold=False,
             italic=False, before=0, after=6, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if text:
        set_font(p.add_run(text), size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    set_font(p.add_run(text), size=11)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size={1: 16, 2: 13, 3: 12}[level],
             color={1: BLUE, 2: BLUE, 3: NAVY}[level], bold=True)
    return p


def add_callout(doc, label, text, fill='F4F6F9', accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + '  ')
    set_font(r, size=10.5, color=accent, bold=True)
    set_font(p.add_run(text), size=10.5, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_phase_table(doc):
    rows = [
        ('Phase 2', 'Network Control & Security', 'Menata dan mengamankan fondasi jaringan melalui managed core, VLAN, firewall, dan kebijakan akses.', 'Segera / prioritas berikutnya'),
        ('Phase 3', 'Central Monitoring & Remote Management', 'Membangun visibilitas terpusat, alert otomatis, backup konfigurasi, dan administrasi jarak jauh yang aman.', 'Rekomendasi masa depan'),
        ('Phase 4', 'Multi-Site Integration', 'Menghubungkan HO, BO, mining site, dan nickel site dengan VPN, routing, redundansi, serta central IT management.', 'Arah jangka panjang'),
    ]
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_widths(table, [1050, 2200, 4300, 1810])
    headers = ['Tahap', 'Fokus', 'Hasil Bisnis/Teknis', 'Prioritas']
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_font(p.add_run(h), size=9.5, color='FFFFFF', bold=True)
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            if i == 0:
                set_cell_shading(cells[i], LIGHT_BLUE)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), size=9.5, color=NAVY if i == 0 else '000000', bold=i == 0)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(4)
    set_font(p.add_run('Executive Summary | Roadmap Infrastruktur Jaringan | 26 Agustus 2026'), size=8.5, color=GRAY)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.85)
section.right_margin = Inches(0.85)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)
add_footer(section)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10
for level, size, color, before, after in [(1, 16, BLUE, 16, 8), (2, 13, BLUE, 12, 6), (3, 12, NAVY, 8, 4)]:
    st = styles[f'Heading {level}']
    st.font.name = 'Calibri'
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

# Masthead
add_para(doc, 'EXECUTIVE SUMMARY', size=10, color=BLUE, bold=True, after=5)
add_para(doc, 'Roadmap Infrastruktur Jaringan\nPhase 2 sampai Phase 4', size=25, color=NAVY, bold=True, after=7)
add_para(doc, 'Network Control, Central Monitoring, dan Integrasi Multi-Site', size=13, color=GRAY, after=16)
for label, value in [('Tujuan', 'Menyediakan arah implementasi jaringan yang aman, terukur, dan scalable.'),
                     ('Cakupan', 'Workshop infrastructure, central monitoring, remote management, dan multi-site integration.'),
                     ('Status', 'Rekomendasi untuk pengambilan keputusan dan perencanaan implementasi.')]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    set_font(p.add_run(label + ': '), size=10.5, color=NAVY, bold=True)
    set_font(p.add_run(value), size=10.5, color='000000')

add_callout(doc, 'REKOMENDASI UTAMA', 'Prioritaskan Phase 2 sebagai fondasi pengendalian dan keamanan jaringan. Setelah standar VLAN, firewall, IP addressing, dan dokumentasi stabil, lanjutkan ke Phase 3 untuk monitoring terpusat, kemudian Phase 4 untuk konektivitas antar-lokasi.', fill='EAF3EE', accent=GREEN)

add_heading(doc, '1. Ringkasan Eksekutif', 1)
add_para(doc, 'Roadmap Phase 2 sampai Phase 4 mengubah jaringan workshop dari infrastruktur yang berfungsi secara lokal menjadi platform yang aman, dapat dipantau, dan siap terintegrasi antar-lokasi. Perubahan ini penting untuk mengurangi downtime, mempercepat penanganan insiden, membatasi risiko keamanan, serta menyediakan dasar bagi ekspansi operasional.')
add_para(doc, 'Investasi sebaiknya dilakukan bertahap. Phase 2 menghasilkan kontrol dasar dan segmentasi. Phase 3 menambahkan kemampuan operasi proaktif melalui monitoring dan remote administration. Phase 4 menghubungkan seluruh site dengan kebijakan keamanan dan manajemen yang konsisten.')

add_heading(doc, '2. Roadmap Keputusan', 1)
add_phase_table(doc)

add_heading(doc, '3. Prioritas Phase 2 - Network Control & Security', 1)
add_para(doc, 'Phase 2 merupakan fondasi teknis yang harus diselesaikan sebelum monitoring dan integrasi multi-site dilakukan. Fokusnya adalah standardisasi arsitektur, pembatasan akses, dan pengurangan ketergantungan pada konfigurasi perangkat yang tidak terdokumentasi.')
add_heading(doc, 'Kapabilitas yang dibangun', 2)
for text in [
    'Managed core network menggunakan switch Layer 2/Layer 3, firewall/router, rack, UPS, dan backbone yang terdokumentasi.',
    'Network segmentation melalui VLAN terpisah untuk user, CCTV, server/NVR, guest Wi-Fi, management, dan perangkat workshop/IoT.',
    'Pemisahan akses antara CCTV, user, dan administrator; akses antar-segment hanya dibuka sesuai kebutuhan bisnis.',
    'Basic security policy: SSH/HTTPS, MFA bila tersedia, password unik, firmware management, backup konfigurasi, dan port unused yang dinonaktifkan.',
]: add_bullet(doc, text)
add_heading(doc, 'Contoh target teknis', 2)
add_para(doc, 'VLAN 20 untuk user, VLAN 30 untuk CCTV, VLAN 40 untuk guest, VLAN 50 untuk server/NVR, VLAN 60 untuk IoT/workshop, dan VLAN 10 untuk management. Guest hanya mendapatkan akses internet; CCTV hanya berkomunikasi dengan NVR; administrator mengakses perangkat melalui management VLAN.')
add_callout(doc, 'KEPUTUSAN YANG DIPERLUKAN', 'Menyetujui standar VLAN, IP addressing, firewall policy, perangkat managed, serta penanggung jawab perubahan konfigurasi.', fill='FFF7E6', accent=GOLD)

add_heading(doc, '4. Prioritas Phase 3 - Central Monitoring & Remote Management', 1)
add_para(doc, 'Phase 3 mengubah operasi jaringan dari reaktif menjadi proaktif. Tim IT memperoleh satu tampilan untuk memantau availability, performa, kesehatan perangkat, kapasitas, dan kejadian keamanan.')
add_heading(doc, 'Kapabilitas yang dibangun', 2)
for text in [
    'Centralized monitoring menggunakan NMS seperti Zabbix, PRTG, LibreNMS, atau platform vendor.',
    'Monitoring SNMPv3, ICMP, Syslog, dan bila diperlukan NetFlow/sFlow untuk status perangkat, bandwidth, error port, suhu, storage NVR, dan client Wi-Fi.',
    'Automatic alerts melalui email, Teams, Slack, Telegram, atau sistem ticketing dengan klasifikasi warning dan critical.',
    'Secure remote administration melalui VPN, MFA, bastion/jump server, role-based access, dan audit log.',
    'Backup konfigurasi otomatis serta change management untuk memastikan setiap perubahan dapat dilacak dan dipulihkan.',
]: add_bullet(doc, text)
add_heading(doc, 'Contoh indikator dashboard', 2)
add_para(doc, 'Availability jaringan 99,8%; perangkat online 48/50; kamera online 15/16; packet loss di bawah 1%; storage NVR di bawah 85%; dan seluruh sesi VPN administrator tercatat.')

add_heading(doc, '5. Prioritas Phase 4 - Multi-Site Integration', 1)
add_para(doc, 'Phase 4 menyediakan konektivitas aman antara Head Office, Branch Office, Mining Site, Nickel Site, dan lokasi lain. Arsitektur ini harus mendukung pertumbuhan jumlah site tanpa mengulang desain dari awal.')
add_heading(doc, 'Kapabilitas yang dibangun', 2)
for text in [
    'Site-to-site VPN berbasis IPsec atau SD-WAN untuk mengamankan komunikasi antar-lokasi.',
    'IP addressing dan routing plan yang konsisten, misalnya 10.1.0.0/16 untuk HO, 10.2.0.0/16 untuk BO, 10.3.0.0/16 untuk Mining, dan 10.4.0.0/16 untuk Nickel Site.',
    'Central IT management untuk firewall, switch, access point, VPN, monitoring, logging, dan firmware.',
    'Redundansi koneksi menggunakan dual ISP, fiber plus 4G/5G, radio link, microwave, atau VSAT sesuai kondisi lokasi.',
    'Pemisahan jaringan IT, OT, CCTV, guest, dan management dengan prinsip least privilege dan Zero Trust.',
]: add_bullet(doc, text)
add_heading(doc, 'Contoh alur insiden', 2)
add_para(doc, 'Saat kamera Mining Site offline, NMS membuat alert, memeriksa status VPN dan switch, lalu administrator masuk melalui VPN dengan MFA. Tim memeriksa PoE, port switch, IP address, packet loss, dan koneksi NVR. Setelah layanan pulih, aktivitas dicatat dalam audit log dan tiket ditutup.')

add_heading(doc, '6. Manfaat Bisnis', 1)
benefits = [
    ('Resiliensi operasional', 'Downtime lebih cepat terdeteksi dan dipulihkan.'),
    ('Keamanan', 'Segmentasi dan kebijakan akses membatasi penyebaran insiden.'),
    ('Efisiensi IT', 'Troubleshooting dan administrasi dapat dilakukan dari pusat.'),
    ('Skalabilitas', 'Site baru dapat mengikuti standar IP, VLAN, VPN, dan monitoring yang sama.'),
    ('Governance', 'Perubahan, log, backup, dan performa memiliki jejak audit yang jelas.'),
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_widths(table, [2500, 6860])
for i, h in enumerate(['Manfaat', 'Dampak yang Diharapkan']):
    set_cell_shading(table.rows[0].cells[i], LIGHT_GRAY)
    set_font(table.rows[0].cells[i].paragraphs[0].add_run(h), size=10, color=NAVY, bold=True)
for a, b in benefits:
    cells = table.add_row().cells
    set_font(cells[0].paragraphs[0].add_run(a), size=10, color=NAVY, bold=True)
    set_font(cells[1].paragraphs[0].add_run(b), size=10)

add_heading(doc, '7. Risiko dan Mitigasi', 1)
risk_rows = [
    ('Perubahan mengganggu operasional', 'Lakukan maintenance window, backup konfigurasi, dan rollback plan.'),
    ('Kapasitas link tidak mencukupi', 'Ukur baseline trafik dan siapkan QoS serta link cadangan.'),
    ('Akses remote disalahgunakan', 'Gunakan VPN, MFA, RBAC, allowlist, dan audit log.'),
    ('Ketergantungan pada satu vendor', 'Dokumentasikan standar terbuka dan evaluasi interoperabilitas.'),
    ('Perangkat lama tidak mendukung segmentasi', 'Migrasikan bertahap dan gunakan firewall/ACL sebagai kontrol sementara.'),
]
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.LEFT
set_table_widths(table, [3000, 6360])
for i, h in enumerate(['Risiko', 'Mitigasi']):
    set_cell_shading(table.rows[0].cells[i], LIGHT_GRAY)
    set_font(table.rows[0].cells[i].paragraphs[0].add_run(h), size=10, color=NAVY, bold=True)
for a, b in risk_rows:
    cells = table.add_row().cells
    set_font(cells[0].paragraphs[0].add_run(a), size=10, color=NAVY, bold=True)
    set_font(cells[1].paragraphs[0].add_run(b), size=10)

add_heading(doc, '8. Rekomendasi Tindak Lanjut', 1)
steps = [
    'Validasi kondisi existing: perangkat, topologi, IP address, kapasitas link, CCTV, Wi-Fi, dan kebutuhan tiap lokasi.',
    'Tetapkan standar Phase 2: VLAN, subnet, firewall rule, naming convention, akses admin, dan format dokumentasi.',
    'Implementasikan pilot pada satu area workshop sebelum roll-out ke seluruh area.',
    'Bangun monitoring Phase 3 setelah perangkat utama memiliki konfigurasi dan identitas yang konsisten.',
    'Susun blueprint Phase 4 setelah kebutuhan antar-site, pilihan link, SLA, dan skenario failover disetujui.',
]
for idx, text in enumerate(steps, 1):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.10
    set_font(p.add_run(text), size=11)

add_callout(doc, 'KESIMPULAN', 'Phase 2 adalah prioritas implementasi karena menjadi fondasi keamanan dan kontrol. Phase 3 memberikan operational excellence melalui monitoring dan remote management. Phase 4 menjadi arah jangka panjang untuk jaringan yang scalable, terintegrasi, dan siap mendukung operasi multi-site.', fill='EAF3EE', accent=GREEN)

doc.core_properties.title = 'Executive Summary Roadmap Infrastruktur Jaringan Phase 2 sampai Phase 4'
doc.core_properties.subject = 'Network Control, Monitoring, Remote Management, dan Multi-Site Integration'
doc.core_properties.author = 'Insting Tekno Sinergi'
doc.core_properties.keywords = 'network, security, monitoring, multi-site, executive summary'
doc.save(OUT)
print(OUT)
